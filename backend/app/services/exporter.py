from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document


@dataclass(frozen=True)
class FlatSection:
    id: int
    parent_id: int | None
    order_index: int
    title: str
    content_md: str


def build_tree(sections: list[FlatSection]) -> dict[int | None, list[FlatSection]]:
    by_parent: dict[int | None, list[FlatSection]] = {}
    for s in sections:
        by_parent.setdefault(s.parent_id, []).append(s)
    for arr in by_parent.values():
        arr.sort(key=lambda x: x.order_index)
    return by_parent


def render_markdown(sections: list[FlatSection]) -> str:
    by_parent = build_tree(sections)
    lines: list[str] = []

    def walk(parent: int | None, depth: int) -> None:
        for sec in by_parent.get(parent, []):
            heading = "#" * min(depth + 2, 6)
            lines.append(f"{heading} {sec.title}".rstrip())
            if sec.content_md.strip():
                lines.append(sec.content_md.strip())
            lines.append("")
            walk(sec.id, depth + 1)

    walk(None, 0)
    return "\n".join(lines).strip() + "\n"




def strip_markdown(value: str) -> str:
    text = value or ""
    text = re.sub(r"^\s{0,3}#{1,6}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s{0,3}>\s?", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*[-*_]{3,}\s*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*[-*+]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*\d+[.)]\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"```[\s\S]*?```", lambda m: m.group(0).replace("```", "").strip(), text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"(\*\*|__)(.*?)\1", r"\2", text)
    text = re.sub(r"(?<!\*)\*(?!\*)([^*]+)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"(?<!_)_(?!_)([^_]+)(?<!_)_(?!_)", r"\1", text)
    text = re.sub(r"~~(.*?)~~", r"\1", text)
    return text.strip()

def export_docx(title: str, sections: list[FlatSection], out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)

    by_parent = build_tree(sections)

    def walk(parent: int | None, depth: int) -> None:
        for sec in by_parent.get(parent, []):
            level = min(depth + 1, 9)
            doc.add_heading(sec.title, level=level)
            for para in strip_markdown(sec.content_md).splitlines():
                if para.strip() == "":
                    continue
                doc.add_paragraph(para)
            walk(sec.id, depth + 1)

    walk(None, 0)
    doc.save(str(out_path))
    return out_path



def _plain_text_lines(sections: list[FlatSection]) -> list[str]:
    by_parent = build_tree(sections)
    lines: list[str] = []

    def walk(parent: int | None, depth: int) -> None:
        for sec in by_parent.get(parent, []):
            indent = '  ' * depth
            lines.append(f"{indent}{sec.title}")
            content = strip_markdown(sec.content_md)
            if content.strip():
                for para in content.splitlines():
                    if para.strip():
                        lines.append(f"{indent}{para.strip()}")
            lines.append('')
            walk(sec.id, depth + 1)

    walk(None, 0)
    return lines


def _wrap_text_by_font(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    if not text:
        return ['']
    lines: list[str] = []
    for raw in text.splitlines() or ['']:
        if not raw.strip():
            lines.append('')
            continue
        current = ''
        for ch in raw:
            trial = current + ch
            bbox = draw.textbbox((0, 0), trial, font=font)
            if bbox[2] - bbox[0] <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = ch
        if current:
            lines.append(current)
    return lines


def export_pdf(title: str, sections: list[FlatSection], out_path: Path) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    font_path = Path(r'C:\Windows\Fonts\msyhl.ttc')
    if font_path.exists():
        title_font = ImageFont.truetype(str(font_path), 28)
        body_font = ImageFont.truetype(str(font_path), 20)
    else:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()

    page_width, page_height = 1240, 1754
    margin_x, margin_y = 80, 80
    max_width = page_width - margin_x * 2
    line_spacing = 12
    title_gap = 24
    body_gap = 8

    base = Image.new('RGB', (page_width, page_height), 'white')
    drawer = ImageDraw.Draw(base)

    plain_lines = [title, ''] + _plain_text_lines(sections)
    wrapped_lines: list[tuple[str, bool]] = []
    for idx, line in enumerate(plain_lines):
        is_title = idx == 0
        font = title_font if is_title else body_font
        for wrapped in _wrap_text_by_font(drawer, line, font, max_width):
            wrapped_lines.append((wrapped, is_title))
            is_title = False

    pages: list[Image.Image] = []
    current = Image.new('RGB', (page_width, page_height), 'white')
    draw = ImageDraw.Draw(current)
    y = margin_y

    def flush_page() -> None:
        nonlocal current, draw, y
        pages.append(current)
        current = Image.new('RGB', (page_width, page_height), 'white')
        draw = ImageDraw.Draw(current)
        y = margin_y

    for line, is_title in wrapped_lines:
        font = title_font if is_title else body_font
        line_height = draw.textbbox((0, 0), line or 'A', font=font)[3] - draw.textbbox((0, 0), line or 'A', font=font)[1]
        needed = line_height + (title_gap if is_title else body_gap)
        if y + needed > page_height - margin_y:
            flush_page()
        draw.text((margin_x, y), line, fill='black', font=font)
        y += needed

    pages.append(current)
    first, *rest = [img.convert('RGB') for img in pages]
    first.save(out_path, save_all=True, append_images=rest)
    return out_path
